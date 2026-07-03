"""
build_presentation.py — DroPhenix scientific presentation (.pptx).
All content is drawn from the project's verified results (validation/run_all.py,
comparison_study.py, demo_unequal_n.py, reference_datasets/). No fabricated data.

Run:  python manuscript/build_presentation.py  ->  manuscript/DroPhenix_presentation.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT = HERE / "DroPhenix_presentation.pptx"

NAVY = RGBColor(0x0D, 0x21, 0x37)
NAVY2 = RGBColor(0x12, 0x2E, 0x47)
GOLD = RGBColor(0xF0, 0xA5, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x33)
SLATE = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xEE, 0xF2, 0xF9)
GREEN = RGBColor(0x00, 0x9E, 0x73)
CARDBG = RGBColor(0xF4, 0xF7, 0xFB)
CARDLINE = RGBColor(0xC5, 0xD7, 0xF0)

HEAD = "Calibri"
BODY = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _tb(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(2)
    return tb, tf


def _run(p, text, size, color, bold=False, italic=False, font=BODY):
    r = p.add_run(); r.text = text; f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font
    f.color.rgb = color; return r


def _fill(shape, color, line=None, lw=0.75):
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line; shape.line.width = Pt(lw)
    shape.shadow.inherit = False


def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH); _fill(r, color)
    slide.shapes._spTree.remove(r._element); slide.shapes._spTree.insert(2, r._element)
    return r


def title_bar(slide, kicker, title):
    _, tf = _tb(slide, Inches(0.6), Inches(0.42), Inches(12.1), Inches(0.4))
    p = tf.paragraphs[0]; _run(p, kicker.upper(), 12, GOLD, bold=True, font=HEAD)
    _, tf = _tb(slide, Inches(0.6), Inches(0.78), Inches(12.1), Inches(0.9))
    p = tf.paragraphs[0]; _run(p, title, 30, NAVY, bold=True, font=HEAD)


def bullets(slide, l, t, w, h, items, size=15, gap=8, color=INK):
    _, tf = _tb(slide, l, t, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        _run(p, "•  ", size, GOLD, bold=True)
        if isinstance(it, tuple):
            _run(p, it[0] + " ", size, color, bold=True); _run(p, it[1], size, color)
        else:
            _run(p, it, size, color)


def card(slide, l, t, w, h, head, body, headcolor=NAVY):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(box, CARDBG, CARDLINE, 1.0)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(10); tf.margin_top = Pt(9); tf.margin_bottom = Pt(8)
    p = tf.paragraphs[0]; p.space_after = Pt(4); _run(p, head, 14, headcolor, bold=True, font=HEAD)
    for i, b in enumerate(body):
        pp = tf.add_paragraph(); pp.line_spacing = 1.03; pp.space_after = Pt(2)
        _run(pp, b, 11.5, INK)
    return box


def stat(slide, l, t, w, number, label, ncolor=GOLD):
    _, tf = _tb(slide, l, t, w, Inches(0.75), MSO_ANCHOR.BOTTOM)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _run(p, number, 40, ncolor, bold=True, font=HEAD)
    _, tf = _tb(slide, l, t + Inches(0.78), w, Inches(0.55))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _run(p, label, 11.5, WHITE)


def content(kicker, title):
    s = prs.slides.add_slide(BLANK); bg(s, WHITE); title_bar(s, kicker, title); return s


def footer(slide, dark=False):
    _, tf = _tb(slide, Inches(0.6), Inches(7.02), Inches(12.1), Inches(0.3))
    p = tf.paragraphs[0]
    _run(p, "DroPhenix Analytics v3.0", 9, GOLD if dark else SLATE, bold=True)
    _run(p, "   ·   drophenix.streamlit.app   ·   SSSIHL", 9, (LIGHT if dark else SLATE))


# ── 1. TITLE ──────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.05), SW, Inches(0.03)); _fill(band, GOLD)
_, tf = _tb(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.4))
_run(tf.paragraphs[0], "SCIENCE FOR SOCIETY  ·  SSSIHL  ·  SAI-NET", 13, GOLD, bold=True, font=HEAD)
_, tf = _tb(s, Inches(0.9), Inches(1.1), Inches(11.6), Inches(1.0))
p = tf.paragraphs[0]; _run(p, "DroPhenix ", 46, WHITE, bold=True, font=HEAD); _run(p, "Analytics", 46, GOLD, bold=True, font=HEAD)
_, tf = _tb(s, Inches(0.9), Inches(2.25), Inches(11.6), Inches(0.6))
_run(tf.paragraphs[0], "A statistically rigorous web platform for multi-assay Drosophila phenotyping and actuarial cross-strain compound ranking", 18, LIGHT, font=HEAD)
for i, (n, lb) in enumerate([("6", "assays"), ("24/24", "tool checks agree"), ("r = 1.0000", "vs gold-standard software"), ("18/18", "reference datasets")]):
    stat(s, Inches(0.9 + i * 3.0), Inches(3.7), Inches(2.7), n, lb)
_, tf = _tb(s, Inches(0.9), Inches(6.35), Inches(11.6), Inches(0.5))
_run(tf.paragraphs[0], "Krishnasalini Gunanathan  ·  Venketesh Sivaramakrishnan   |   Dept. of Biosciences, SSSIHL", 13, LIGHT)
footer(s, dark=True)

# ── 2. THE PROBLEM ────────────────────────────────────────────────────────────
s = content("Motivation", "The problem we set out to solve")
bullets(s, Inches(0.6), Inches(1.75), Inches(6.4), Inches(5),
        [("Fragmented analysis.", "Six standard Drosophila assays — climbing, lifespan, Y-maze, sleep/circadian (DAMS), larval crawling, grooming — are analysed in disconnected spreadsheets and single-purpose scripts."),
         ("Under-powered statistics.", "Climbing reduced to one endpoint %; lifespan reported as median only, omitting Kaplan–Meier, log-rank and Cox; periodograms often uncorrected for multiplicity."),
         ("No unequal-N method.", "When one strain has 3 vials and another 7, averaging group means is biased toward the smaller, noisier strain — yet this is exactly the screening situation."),
         ("Verifiability gap.", "Results are rarely benchmarked against published data or gold-standard software.")], size=14.5, gap=12)
card(s, Inches(7.4), Inches(1.9), Inches(5.3), Inches(4.4), "The unmet need",
     ["One reproducible workflow that:", "• unifies all six assays,",
      "• applies correct, assumption-aware statistics,", "• handles unequal-N strain comparison rigorously,",
      "• and can be independently verified —", "  not merely trusted.", "",
      "DroPhenix was built to meet each of these."], headcolor=GOLD)
footer(s)

# ── 3. WHAT WE BUILT ──────────────────────────────────────────────────────────
s = content("Overview", "What we developed: six assays + cross-strain ranking")
mods = [("Climbing", "Negative geotaxis: PI, AUC, t50; video fly-counting"),
        ("Lifespan", "Kaplan–Meier, log-rank, Cox proportional hazards"),
        ("Y-maze", "Spontaneous alternation; turn-bias / handedness"),
        ("DAMS", "Sleep (5-min rule) + circadian period (periodogram)"),
        ("Larval crawling", "Video tracking: speed, path, tortuosity"),
        ("Grooming", "% time grooming, bout structure"),
        ("Motor Decline Rate", "Exponential decline of PI over age"),
        ("Trajectory clustering", "DTW + Ward unsupervised phenotype grouping"),
        ("Drug Ranking", "Cross-strain, cross-assay meta-analytic score")]
x0, y0, w, h, gx, gy = Inches(0.6), Inches(1.7), Inches(4.0), Inches(1.5), Inches(0.15), Inches(0.16)
for i, (nm, ds) in enumerate(mods):
    r, c = divmod(i, 3)
    card(s, x0 + c * (w + gx), y0 + r * (h + gy), w, h, nm, [ds])
footer(s)

# ── 4. THE INNOVATION ─────────────────────────────────────────────────────────
s = content("Core innovation", "Actuarial cross-strain compound comparison under unequal N")
bullets(s, Inches(0.6), Inches(1.7), Inches(6.3), Inches(4.6),
        [("Per strain.", "Drug-vs-vehicle effect = Hedges' g, whose variance depends on the two group sizes — so unequal N is handled at source."),
         ("Across strains.", "Combine by inverse-variance / DerSimonian–Laird random-effects pooling: larger strains get more weight; small ones are down-weighted, not discarded."),
         ("Across assays.", "Pool into one Therapeutic Score, oriented so positive = benefit, with heterogeneity (I²)."),
         ("Why it matters.", "This is the minimum-variance, correctly-calibrated way to rank a compound across unbalanced disease models — the first question a reviewer asks of a screen.")], size=14.5, gap=11)
if (FIG / "Figure2_innovation.png").exists():
    s.shapes.add_picture(str(FIG / "Figure2_innovation.png"), Inches(7.05), Inches(1.75), width=Inches(5.75))
_, tf = _tb(s, Inches(7.05), Inches(6.35), Inches(5.75), Inches(0.5))
_run(tf.paragraphs[0], "Simulation (3000 runs): inverse-variance pooling gives 19.9% lower RMSE and 95.8% CI coverage vs naive averaging.", 11, SLATE, italic=True)
footer(s)

# ── 5. MODULES & CALCULATIONS ─────────────────────────────────────────────────
s = content("Methods", "The calculation behind each module — and why")
rows = [("Climbing", "PI = 2·(climbed/total)·100 − 100", "Standard negative-geotaxis readout"),
        ("Lifespan", "KM Ŝ(t)=Π(1−d/n); log-rank χ²; Cox HR=e^β", "Actuarial standard for censored data"),
        ("Circadian", "Sokolove–Bushell Q_P periodogram (Bonferroni)", "Field standard; controls false positives"),
        ("Sleep", "bout = ≥5 min zero activity", "Operational Drosophila sleep definition"),
        ("Y-maze", "%alt=(triads/(entries−2)); bias=R/(R+L)", "Working memory + behavioural individuality"),
        ("MDR", "PI(age)=PI₀·e^(−k·age)", "Decline is age-dependent, not per-second"),
        ("Ranking", "Hedges' g → inverse-variance / DL pooling", "Unequal-N, minimum-variance combination")]
_, tf = _tb(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(0.3))
pp = tf.paragraphs[0]
for txt, wd in [("Module", 2.3), ("Calculation", 6.0), ("Why we use it", 3.8)]:
    _run(pp, txt.ljust(0), 13, GOLD, bold=True); _run(pp, "\t", 13, GOLD)
y = Inches(2.05)
for nm, fm, why in rows:
    _, tf = _tb(s, Inches(0.6), y, Inches(2.3), Inches(0.62), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], nm, 12.5, NAVY, bold=True)
    _, tf = _tb(s, Inches(2.95), y, Inches(5.7), Inches(0.62), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], fm, 12, INK, font="Cambria")
    _, tf = _tb(s, Inches(8.85), y, Inches(3.9), Inches(0.62), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], why, 11.5, SLATE)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y + Inches(0.6), Inches(12.15), Pt(0.75)); _fill(ln, CARDLINE)
    y += Inches(0.66)
footer(s)

# ── 6. STATISTICAL ENGINE ─────────────────────────────────────────────────────
s = content("Rigour", "The statistics engine: assumption-aware by design")
bullets(s, Inches(0.6), Inches(1.75), Inches(6.2), Inches(5),
        [("Test selection is gated, not defaulted.", "Shapiro–Wilk (normality) + Levene (variance) decide between Welch t / ANOVA and Mann–Whitney / Kruskal–Wallis."),
         ("Effect sizes are mandatory.", "Cohen's d and small-sample-corrected Hedges' g; rank-biserial for non-parametric — significance is never reported alone."),
         ("Multiplicity is controlled.", "Benjamini–Hochberg false-discovery-rate correction on all pairwise post-hoc tests."),
         ("Uncertainty is quantified.", "Percentile bootstrap 95% confidence intervals (fixed seed → reproducible).")], size=14.5, gap=13)
card(s, Inches(7.15), Inches(1.9), Inches(5.55), Inches(4.4), "Why this matters",
     ["Behavioural screens are small and noisy.",
      "Defaulting to ANOVA on non-normal, unequal-variance, unequal-N data inflates false positives.",
      "",
      "Gating the test on its assumptions, forcing effect sizes, and controlling the FDR is the",
      "difference between a result that survives peer review and one that does not."], headcolor=GOLD)
footer(s)

# ── 7. TOOLS & TECHNIQUES ─────────────────────────────────────────────────────
s = content("Engineering", "Tools and techniques — and why we chose them")
tools = [("Streamlit", "Zero-install browser UI; no login"),
         ("NumPy / pandas / SciPy", "Vectorised numerics, curve-fitting, Lomb–Scargle"),
         ("statsmodels", "Peer-reviewed survival core (KM, log-rank, Cox) + meta-analysis"),
         ("scikit-learn", "PCA, silhouette, StandardScaler for clustering"),
         ("OpenCV", "Otsu-threshold fly/larva detection (no training needed)"),
         ("Matplotlib / Plotly", "300-dpi publication figures; interactive plots")]
x0, y0, w, h, gx, gy = Inches(0.6), Inches(1.7), Inches(6.05), Inches(1.2), Inches(0.2), Inches(0.18)
for i, (nm, ds) in enumerate(tools):
    r, c = divmod(i, 2)
    card(s, x0 + c * (w + gx), y0 + r * (h + gy), w, h, nm, [ds])
_, tf = _tb(s, Inches(0.6), Inches(6.25), Inches(12.1), Inches(0.5))
p = tf.paragraphs[0]
_run(p, "Deterministic, not a trained model: ", 13, NAVY, bold=True)
_run(p, "detection uses Otsu thresholding, clustering is unsupervised (DTW) — no black-box weights, every step is inspectable.", 13, INK)
footer(s)

# ── 8. DATA SOURCES ───────────────────────────────────────────────────────────
s = content("Data", "Data sources — real published anchors, honestly labelled")
bullets(s, Inches(0.6), Inches(1.7), Inches(6.4), Inches(5),
        [("Real raw data.", "Freireich et al. 1963 6-MP leukemia trial (survival); PlantGrowth / Dobson 1983 (ANOVA)."),
         ("Real published parameters.", "Konopka & Benzer 1971 free-running periods (24 / 19 / 28 h)."),
         ("Text-verified literature values.", "Alternation chance 50% (Lewis 2017); turn bias 0.50 & 23.5% extreme (de Bivort 2015); WT lifespan ≈50 d (Linford 2013); grooming ≈13% (Qiao 2018); larval ≈1 mm/s."),
         ("Reference datasets.", "≥3 per assay in template format (reference_datasets/), each cited; simulations clearly labelled — no dataset is presented as raw data we did not have.")], size=14, gap=11)
card(s, Inches(7.4), Inches(1.9), Inches(5.3), Inches(4.4), "Verification-first principle",
     ["Where machine-readable raw supplements were",
      "available, we used them directly.",
      "Where they were not, we simulated from",
      "explicitly cited published parameters and",
      "labelled them as such.", "",
      "18/18 reference datasets reproduce their", "cited published expectation."], headcolor=GOLD)
footer(s)

# ── 9. VALIDATION / TESTING PARAMETERS ────────────────────────────────────────
s = content("Validation", "Testing parameters — what we chose and why")
prows = [("Sleep threshold", "≥5 min inactivity", "Hendricks/Shaw 2000 operational definition"),
         ("Alternation chance", "50%", "Lewis 2017; emerges from the arm geometry"),
         ("Reference periods", "24 / 19 / 28 h", "Konopka & Benzer 1971 wild-type / perS / perL"),
         ("Periodogram range", "16–32 h, Bonferroni", "Circadian band; controls family-wise error"),
         ("Turn-bias spread", "23.5% outside [0.3,0.7]", "de Bivort 2015 Canton-S distribution"),
         ("MDR axis", "Age (days)", "Negative geotaxis is age-dependent (Gargano 2005)"),
         ("Ties method (Cox)", "Breslow", "Matches published Klein–Moeschberger value")]
_, tf = _tb(s, Inches(0.6), Inches(1.62), Inches(12.1), Inches(0.3))
_run(tf.paragraphs[0], "Every parameter is set from the literature, not tuned to flatter the result.", 13.5, SLATE, italic=True)
y = Inches(2.1)
for nm, val, why in prows:
    _, tf = _tb(s, Inches(0.6), y, Inches(3.0), Inches(0.6), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], nm, 12.5, NAVY, bold=True)
    _, tf = _tb(s, Inches(3.7), y, Inches(3.1), Inches(0.6), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], val, 12.5, GREEN, bold=True, font="Cambria")
    _, tf = _tb(s, Inches(6.95), y, Inches(5.8), Inches(0.6), MSO_ANCHOR.MIDDLE)
    _run(tf.paragraphs[0], why, 11.5, INK)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y + Inches(0.58), Inches(12.15), Pt(0.75)); _fill(ln, CARDLINE)
    y += Inches(0.63)
footer(s)

# ── 10. RESULTS: VALIDATION ───────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
_, tf = _tb(s, Inches(0.6), Inches(0.5), Inches(12.1), Inches(0.4)); _run(tf.paragraphs[0], "RESULTS · VALIDATION", 12, GOLD, bold=True, font=HEAD)
_, tf = _tb(s, Inches(0.6), Inches(0.85), Inches(12.1), Inches(0.7)); _run(tf.paragraphs[0], "Every module reproduces published benchmarks", 30, WHITE, bold=True, font=HEAD)
for i, (n, lb) in enumerate([("12/12", "validation suites pass"), ("χ² = 16.79", "Freireich log-rank (published)"), ("HR = 4.52", "Cox hazard ratio (= published)"), ("r = 0.989", "MDR k-recovery")]):
    stat(s, Inches(0.7 + i * 3.05), Inches(1.9), Inches(2.8), n, lb)
if (FIG / "Figure1_validation.png").exists():
    s.shapes.add_picture(str(FIG / "Figure1_validation.png"), Inches(1.35), Inches(3.5), width=Inches(10.6))
_, tf = _tb(s, Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.4))
_run(tf.paragraphs[0], "Kaplan–Meier vs lifelines · climbing kinetics · MDR over age · periodogram · handedness · drug ranking", 11, LIGHT, italic=True)

# ── 11. RESULTS: COMPARISON ───────────────────────────────────────────────────
s = content("Results · comparison", "Agreement with seven established, independent tools")
for i, (n, lb) in enumerate([("24/24", "metrics agree"), ("r = 1.0000", "DroPhenix vs reference"), ("7", "reference tools")]):
    st_x = Inches(0.7 + i * 2.15)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, st_x, Inches(1.75), Inches(1.95), Inches(1.5)); _fill(box, NAVY)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _run(p, n, 26, GOLD, bold=True, font=HEAD)
    pp = tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER; _run(pp, lb, 11, WHITE)
bullets(s, Inches(0.7), Inches(3.55), Inches(6.0), Inches(3.2),
        [("R survival + base R", "— KM, log-rank, Cox (Breslow), ANOVA, Kruskal: exact"),
         ("lifelines / statsmodels", "— survival & meta-analysis: exact to 6 dp"),
         ("scipy / pingouin", "— ANOVA / Kruskal / Hedges g: exact"),
         ("astropy + pyBOAT", "— circadian period 24.00 vs 23.96 h (two methods)")], size=13.5, gap=9)
if (FIG / "Figure3_tool_agreement.png").exists():
    s.shapes.add_picture(str(FIG / "Figure3_tool_agreement.png"), Inches(7.2), Inches(1.65), height=Inches(4.9))
footer(s)

# ── 12. WHY RELIABLE ──────────────────────────────────────────────────────────
s = content("Trust", "Why DroPhenix is reliable")
items = [("Reproduces published data", "Textbook datasets (Freireich, PlantGrowth, Konopka–Benzer) return their exact published statistics."),
         ("Agrees with gold-standard software", "24/24 metrics match R survival, lifelines, statsmodels, scipy, pingouin, astropy, pyBOAT (r = 1.0000)."),
         ("Correct by construction", "Assumption-gated tests, FDR control, the right measurement axis (e.g. MDR over age)."),
         ("Independently verifiable", "The full validation + comparison suite ships; every number regenerates with one command."),
         ("Honest about limits", "Method notes (Breslow vs Efron), un-buildable tools, and simulated-vs-real data are all stated explicitly.")]
x0, y0, w, h, gy = Inches(0.6), Inches(1.7), Inches(12.15), Inches(0.92), Inches(0.12)
for i, (hd, bd) in enumerate(items):
    b = card(s, x0, y0 + i * (h + gy), w, h, hd, [bd])
footer(s)

# ── 13. APPLICATIONS ──────────────────────────────────────────────────────────
s = content("Impact", "Applications — how it helps")
apps = [("Neurodegeneration screening", "Quantify a compound's effect in ALS / PD / HD fly models across climbing, lifespan, cognition, sleep and locomotion."),
        ("Cross-strain compound triage", "Rank candidates fairly across disease backgrounds of unequal size — the core novelty."),
        ("Translational scoring", "Prioritise broad, safe, sustained rescue over high-efficacy but sex-discordant or WT-perturbing hits."),
        ("Ageing & circadian biology", "Age-resolved motor decline, sleep architecture and free-running period in one workflow."),
        ("Teaching & reproducibility", "A transparent, benchmarked reference implementation of standard behavioural statistics."),
        ("Open science", "Free, no-login web app + open code + data templates lower the barrier for any fly lab.")]
x0, y0, w, h, gx, gy = Inches(0.6), Inches(1.7), Inches(6.05), Inches(1.5), Inches(0.2), Inches(0.16)
for i, (nm, ds) in enumerate(apps):
    r, c = divmod(i, 2)
    card(s, x0 + c * (w + gx), y0 + r * (h + gy), w, h, nm, [ds])
footer(s)

# ── 14. CLOSING ───────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.4), SW, Inches(0.03)); _fill(band, GOLD)
_, tf = _tb(s, Inches(0.9), Inches(1.1), Inches(11.5), Inches(1.2))
_run(tf.paragraphs[0], "From fragmented spreadsheets to one verifiable workflow", 34, WHITE, bold=True, font=HEAD)
bullets(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(3),
        [("Six assays, one platform,", "with publication-grade statistics throughout."),
         ("A genuinely novel innovation:", "actuarial, unequal-N cross-strain compound ranking (19.9% more efficient than naive averaging)."),
         ("Validated and benchmarked:", "reproduces published data and agrees with gold-standard software on 24/24 checks."),
         ("Free and open:", "drophenix.streamlit.app · GitHub · GPL-3 · manuscript prepared for Nucleic Acids Research.")],
        size=15, gap=11, color=LIGHT)
footer(s, dark=True)

prs.save(str(OUT))
print("wrote", OUT, "|", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
