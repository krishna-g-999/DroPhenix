"""
DroPhenix — dam_report.py
A Word report from a `dam_pipeline.DamRun` (or two), and the check that every
number printed in it traces back to the run's tables.

The prose is generic -- it describes what was found in terms of the design's
own level names -- and it states what the design cannot support (confounded
contrasts, unpaired timepoints, differing photoperiods) as plainly as what it
can. Every table cell is a rounded copy of a value in the run's per-fly or
contrast frames; `verify()` reads the .docx back and confirms that.

    build(run, path, run_b=None, cmp=None, figures=...)  -> path
    verify(path, run, run_b=None, cmp=None)              -> (ok, report)
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from modules import dam_pipeline as P

DP = {m: 0 for m in P.LABEL}
DP["Waking_activity"] = 2
for m in ("Day_sleep_per_light_h", "Night_sleep_per_dark_h", "Day_activity_per_light_h", "Night_activity_per_dark_h"):
    DP[m] = 1


def fmt(x, dp: int) -> str:
    """Half-up rounding as a string; NaN -> em dash."""
    try:
        x = float(x)
    except (TypeError, ValueError):
        return "—"
    if not np.isfinite(x):
        return "—"
    q = 10 ** dp
    v = np.floor(abs(x) * q + 0.5) / q
    return f"{-v if x < 0 else v:.{dp}f}"


def sgn(x, dp: int) -> str:
    s = fmt(x, dp)
    return s if s.startswith("-") or s == "—" else "+" + s


def pq(x) -> str:
    try:
        x = float(x)
    except (TypeError, ValueError):
        return "—"
    return "—" if not np.isfinite(x) else (fmt(x, 4) if x >= 1e-4 else "<0.0001")


# ─────────────────────────────────────────────────────────────────────────────
# Display tables (strings), shared by build() and verify()
# ─────────────────────────────────────────────────────────────────────────────
def group_rows(run: P.DamRun) -> Tuple[List[str], List[List[str]]]:
    d = run.design; pf = run.per_fly
    keys = [d.genotype_col] + ([d.stratify_col] if d.stratify_col else []) + [d.treatment_col]
    groups = list(pf.groupby(keys).groups.keys())
    head = ["Parameter"] + [" ".join(str(x) for x in (g if isinstance(g, tuple) else (g,))) for g in groups]
    rows = []
    for m in P.METRICS:
        r = [f"{P.LABEL[m]} ({P.UNIT[m]})"]
        for g in groups:
            sel = np.ones(len(pf), bool)
            for k, v in zip(keys, g if isinstance(g, tuple) else (g,)):
                sel &= (pf[k] == v).to_numpy()
            v = pf.loc[sel, m].dropna()
            if len(v) == 0:
                r.append("—"); continue
            sem = v.std(ddof=1) / np.sqrt(len(v)) if len(v) > 1 else np.nan
            r.append(f"{fmt(v.mean(), DP[m])}±{fmt(sem, DP[m])}")
        rows.append(r)
    rows.append(["n flies"] + [str(int(len(pf.groupby(keys).groups[g]))) for g in groups])
    return head, rows


def contrast_rows(df: pd.DataFrame, filt=None, params: Optional[Sequence[str]] = None) -> List[dict]:
    if df is None or df.empty:
        return []
    d = df[df["analysis"] == "primary"]
    if filt is not None:
        d = d[filt(d)]
    order = params or (P.METRICS + [p for p, _, _ in P.PHASE_PARAMS])
    rank = {m: i for i, m in enumerate(order)}
    d = d[d["metric"].isin(rank)].copy()
    d["_k"] = d["metric"].map(rank)
    d = d.sort_values("_k", kind="stable")   # stable: several contrasts may share a metric
    out = []
    for _, r in d.iterrows():
        m = r["metric"]; dp = DP.get(m, 0)

        # A contrast in which no stratum had two usable flies per arm carries
        # the test's identity and nothing else: `stratified()` returns early,
        # so mean_ref, diff, pct_change, g and I2 are all absent. Printing
        # "n/a" is the honest row; indexing the missing columns used to raise
        # KeyError and take the whole report build down with it, which is how
        # a design with one empty cell lost its entire document.
        def g(col):
            return r[col] if col in r.index else np.nan

        pc = pd.to_numeric(g("pct_change"), errors="coerce")
        diff = g("diff")
        out.append({
            "metric": P.LABEL.get(m, m), "n": str(g("n_per_stratum")),
            "ref": fmt(g("mean_ref"), dp), "cmp": fmt(g("mean_cmp"), dp),
            "change": (f"{fmt(diff, dp)} ({sgn(pc, 0)}%)" if pd.notna(pc) else fmt(diff, dp)),
            "g": (f"{sgn(g('hedges_g'), 2)} [{sgn(g('g_lo'), 2)}, {sgn(g('g_hi'), 2)}]"
                  if pd.notna(g("hedges_g")) else "n/a"),
            "I2": f"{fmt(g('I2_pct'), 0)}%" if pd.notna(g("I2_pct")) else "n/a",
            "p": pq(g("p")), "q": pq(g("q_BH")),
            "sig": bool(g("significant")) if pd.notna(g("significant")) else False,
            "clean": bool(g("monitor_clean")) if pd.notna(g("monitor_clean")) else False,
        })
    return out


def outlier_rows(run: P.DamRun) -> List[List[str]]:
    d = run.design; pf = run.per_fly[run.per_fly["Outlier"]]
    rows = []
    for _, r in pf.iterrows():
        rows.append([str(r["Monitor"]), str(r["Channel"]),
                     " ".join(str(r[k]) for k in (d.genotype_col, d.stratify_col, d.treatment_col) if k),
                     str(r["Outlier_reason"]),
                     fmt(r["Total_sleep_min_per_day"], 0), sgn(r.get("z_Total_sleep_min_per_day", np.nan), 1),
                     fmt(r["Total_activity_per_day"], 0), sgn(r.get("z_Total_activity_per_day", np.nan), 1)])
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Build
# ─────────────────────────────────────────────────────────────────────────────
def build(run: P.DamRun, path: str, run_b: Optional[P.DamRun] = None, cmp: Optional[dict] = None,
          figures: Optional[Dict[str, bytes]] = None, title: Optional[str] = None) -> str:
    """Write the .docx. `figures` maps caption -> PNG bytes (in order)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    d = run.design
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.9))

    def H(t, lvl=1):
        doc.add_heading(t, level=lvl)

    def para(t, bold_lead: Optional[str] = None, italic=False, size=None, colour=None):
        p = doc.add_paragraph()
        if bold_lead:
            r = p.add_run(bold_lead + " "); r.bold = True
        r = p.add_run(t); r.italic = italic
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor(*colour)
        return p

    def table(head: Sequence[str], rows: Sequence[Sequence[str]], hit: Optional[Sequence[bool]] = None,
              near: Optional[Sequence[bool]] = None, size=8):
        t = doc.add_table(rows=1, cols=len(head)); t.style = "Light Grid Accent 1"
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
                if i > 0:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if hit is not None and hit[ri]:
                    r.bold = True; r.font.color.rgb = RGBColor(0x0E, 0x6B, 0x45)
                elif near is not None and near[ri]:
                    r.font.color.rgb = RGBColor(0x9A, 0x5E, 0x00)
        doc.add_paragraph()
        return t

    def contrast_table(rows: List[dict], ref_label: str, cmp_label: str):
        head = ["Parameter", "n per stratum", ref_label, cmp_label, "Change", "Hedges g [95% CI]", "I²", "p", "q (BH)"]
        body = [[r["metric"], r["n"], r["ref"], r["cmp"], r["change"], r["g"], r["I2"], r["p"],
                 r["q"] + (" **" if r["sig"] else "")] for r in rows]
        table(head, body, hit=[r["sig"] for r in rows],
              near=[(not r["sig"]) and r["q"] not in ("—",) and r["q"][0] != "<" and float(r["q"]) < 0.10 for r in rows])

    runs = [run] + ([run_b] if run_b else [])
    ttl = title or ("DAM sleep and locomotor analysis: " + " and ".join(r.name for r in runs))
    doc.add_heading(ttl, level=0)
    para(f"Generated by DroPhenix dam_pipeline. Sleep = ≥5 min of zero beam crossings (Shaw et al. 2000; "
         f"Hendricks et al. 2000). Parameters follow Pfeiffenberger et al. (2010), Chiu et al. (2010) and ShinyR-DAM "
         f"(Cichewicz & Hirsh 2018). Every number in this document is a rounded copy of a value in the analysis "
         f"tables and is checked against them by dam_report.verify().", size=9, colour=(0x50, 0x58, 0x64))

    # ── Summary ──
    H("Summary")
    rows = []
    for r in runs:
        s = r.summary()
        for fam, v in s["families"].items():
            rows.append([r.name, fam.replace("_", " "), str(v["n_tests"]), str(v["n_significant"])])
    if cmp is not None:
        p = cmp["timepoint"]; p = p[p["analysis"] == "primary"]
        rows.append([f"{run.name} → {run_b.name}", "timepoint", str(len(p)), str(int(p["significant"].sum()))])
    table(["Run", "Family", "Primary tests", "Significant (FDR 5%)"], rows, size=9)

    # ── Recording ──
    H("1. What the recordings contain")
    head = [""] + [r.name for r in runs]
    q = [r.quality for r in runs]
    rows = [["Recording bin (min)"] + [fmt(x["bin_min"], 0) for x in q],
            ["Photoperiod, measured from the light channel"] + [f"{r.photoperiod_min/60:.1f}L:{24 - r.photoperiod_min/60:.1f}D" for r in runs],
            ["Usable complete days"] + [str(x["n_usable_days"]) for x in q],
            ["Analysis window"] + [f"{r.window[0].date()} to {(r.window[1] - pd.Timedelta(1, "D")).date()}" if r.window else "—" for r in runs],
            ["Monitors' light channels agree"] + [f"{100*x['agreement']:.1f}%" for x in q],
            ["Light channel used for ZT"] + [str(x["cycling"]) for x in q],
            ["Flies analysed"] + [str(len(r.per_fly)) for r in runs],
            ["Flagged outliers"] + [str(int(r.per_fly["Outlier"].sum())) for r in runs]]
    table(head, rows, size=9)
    for r in runs:
        for f in r.quality["faults"]:
            para(f, bold_lead=f"{r.name}:", size=9)

    # ── Design ──
    H("2. What the design can answer")
    para(f"Reference genotype: {d.ref_genotype}. Reference treatment: {d.ref_treatment}. "
         f"Stratum: {d.stratify_col or 'none'}. A contrast is monitor-clean when both of its arms sit in the same "
         f"set of monitors; contrasts that cross a monitor boundary are reported but marked confounded, because "
         f"instrument and biology cannot then be separated. Circadian period is not reported: the recordings are "
         f"in light–dark, where a periodogram measures the lamp timer.")
    dt = P.design_table(run.per_fly, d)
    table(list(dt.columns), dt.astype(str).values.tolist(), size=9)
    if cmp is not None:
        pr = cmp["pairing"]
        para(("The tubes are NOT the same individuals across runs (within-group correlation of each tube's value "
              f"between runs: {', '.join(f'{r.metric} rho={r.within_group_rho:+.2f} (p={r.p:.2f})' for r in pr.itertuples())}). "
              "The timepoint comparison is therefore unpaired.") if not pr.attrs.get("paired") else
             "The tubes carry the same individuals across runs; a paired analysis would be more powerful than the one below.",
             size=9)

    # ── Outliers ──
    H("3. Outliers")
    para(f"Pre-specified rule: modified z (Iglewicz & Hoaglin 1993) of the deviation from the fly's own "
         f"{d.genotype_col} × {d.stratify_col or ''} × {d.treatment_col} group median, MAD pooled per run, "
         f"|z| > {P.OUTLIER_Z} on total sleep or total activity. The analysis with every fly is primary; the "
         f"analysis without the flagged flies is a sensitivity check and is reported beside it.", size=9)
    for r in runs:
        orows = outlier_rows(r)
        if orows:
            para(r.name, bold_lead="Flagged in", size=9)
            table(["Monitor", "Channel", "Group", "Flagged on", "Total sleep", "z", "Total activity", "z"], orows, size=8)
        else:
            para(f"{r.name}: no fly flagged.", size=9)

    # ── Group tables ──
    H("4. Every parameter by group")
    for r in runs:
        para(f"{r.name}: mean ± SEM over {r.quality['n_usable_days']} complete days; all flies.", size=9)
        head, rows = group_rows(r)
        table(head, rows, size=7)

    # ── Genotype ──
    H("5. Genotype contrasts")
    para(f"Among {d.ref_treatment} flies. Primary test: van Elteren's stratified Wilcoxon with {d.stratify_col or 'no'} "
         f"stratum, exact by enumeration; effect: fixed-effect pooled Hedges' g; Benjamini–Hochberg across the "
         f"primary tests of the family. Green rows survive FDR at 5%; amber rows sit at q < 0.10. "
         f"'**' marks q < 0.05.", size=9)
    for r in runs:
        gen = r.contrasts["genotype"]
        for (ref, cmp_), sub in gen[gen["analysis"] == "primary"].groupby(["ref", "cmp"], sort=False):
            clean = bool(sub["monitor_clean"].iloc[0])
            H(f"{r.name}: {ref} vs {cmp_}" + ("" if clean else " (confounded with monitor)"), 2)
            contrast_table(contrast_rows(gen, lambda x, a=ref, b=cmp_: (x["ref"] == a) & (x["cmp"] == b)), ref, cmp_)
        wm = r.contrasts.get("genotype_within_monitor")
        if wm is not None and not wm.empty:
            H(f"{r.name}: genotype contrasts inside one monitor", 2)
            para("Both genotypes in the same monitor at the same treatment: the check that separates genotype "
                 "from the instrument and from channel position.", size=9)
            for (strat, ref, cmp_), sub in wm[wm["analysis"] == "primary"].groupby(["stratum", "ref", "cmp"], sort=False):
                para(f"{strat}: {ref} vs {cmp_}", bold_lead="", size=9)
                contrast_table(contrast_rows(wm, lambda x, s=strat, a=ref, b=cmp_: (x["stratum"] == s) & (x["ref"] == a) & (x["cmp"] == b)), ref, cmp_)

    # ── Treatment ──
    H("6. Treatment contrasts")
    for r in runs:
        for key, lab in (("treatment_pooled_clean", "pooled over monitor-clean blocks (primary)"),
                         ("treatment_pooled_all", "pooled over all blocks")):
            t = r.contrasts.get(key)
            if t is None or t.empty:
                continue
            for trt in sorted(t["cmp"].unique()):
                H(f"{r.name}: {trt} vs {d.ref_treatment}, {lab}", 2)
                contrast_table(contrast_rows(t, lambda x, c=trt: x["cmp"] == c), d.ref_treatment, trt)
            tn = r.contrasts_noout.get(key)
            if key == "treatment_pooled_clean" and tn is not None and not tn.empty:
                for trt in sorted(tn["cmp"].unique()):
                    H(f"{r.name}: {trt} vs {d.ref_treatment}, monitor-clean blocks, outliers removed", 2)
                    contrast_table(contrast_rows(tn, lambda x, c=trt: x["cmp"] == c), d.ref_treatment, trt)
        tw = r.contrasts.get("treatment")
        if tw is not None and not tw.empty:
            H(f"{r.name}: treatment within each genotype", 2)
            for (strat, trt), sub in tw[tw["analysis"] == "primary"].groupby(["stratum", "cmp"], sort=False):
                clean = bool(sub["monitor_clean"].iloc[0])
                para(f"{strat.replace(' stratified', '')}: {trt} vs {d.ref_treatment}" + ("" if clean else " (confounded with monitor)"), size=9)
                contrast_table(contrast_rows(tw, lambda x, s=strat, c=trt: (x["stratum"] == s) & (x["cmp"] == c)), d.ref_treatment, trt)

    # ── Timepoint ──
    if cmp is not None:
        H(f"7. {run.name} → {run_b.name}, within each group")
        La, Lb = run.photoperiod_min / 60, run_b.photoperiod_min / 60
        if abs(La - Lb) > 0.1:
            para(f"The two runs were on different photoperiods ({La:.1f} h vs {Lb:.1f} h of light). Raw day and night "
                 f"sleep therefore move even if the fly does not; day and night sleep and activity are also given per "
                 f"hour of the phase actually recorded.", size=9)
        tp = cmp["timepoint"]
        for strat, sub in tp[tp["analysis"] == "primary"].groupby("stratum", sort=False):
            H(strat.replace(" stratified", ""), 2)
            contrast_table(contrast_rows(tp, lambda x, s=strat: x["stratum"] == s), run.name, run_b.name)

    # ── Sensitivity ──
    H("8. Sensitivity to the outliers")
    rows = []
    for r in runs:
        for key in ("genotype", "treatment_pooled_clean"):
            a, b = r.contrasts.get(key), r.contrasts_noout.get(key)
            if a is None or b is None or a.empty or b.empty:
                continue
            a = a[a["analysis"] == "primary"]; b = b[b["analysis"] == "primary"]
            m = a.merge(b, on=["metric", "stratum", "ref", "cmp"], suffixes=("_a", "_b"))
            for _, x in m.iterrows():
                if x["metric"] not in ("Total_sleep_min_per_day", "Night_sleep_min_per_day", "Longest_bout_min", "Total_activity_per_day"):
                    continue
                rows.append([r.name, f"{x['ref']} vs {x['cmp']} ({x['stratum']})", P.LABEL[x["metric"]],
                             sgn(x["hedges_g_a"], 2), pq(x["q_BH_a"]) + (" **" if x["significant_a"] else ""),
                             sgn(x["hedges_g_b"], 2), pq(x["q_BH_b"]) + (" **" if x["significant_b"] else "")])
    table(["Run", "Contrast", "Parameter", "g, all flies", "q, all flies", "g, no outliers", "q, no outliers"], rows, size=8)

    # ── Methods ──
    H("9. Methods")
    for lead, txt in [
        ("Sleep.", "≥5 min of continuous zero beam crossings."),
        ("Zeitgeber time.", "ZT0 = lights-on, from the light channel of the monitor whose channel actually cycles; light spans shorter than 60 min are ignored as flicker."),
        ("Photoperiod.", "Median of complete lights-on → lights-off intervals; never assumed."),
        ("Window.", "The longest run of complete calendar days on which the light schedule matched the run's usual schedule, with no gap over 30 min and no light-channel flicker."),
        ("Deaths.", "Tubes with fewer than 50 beam crossings in the window are excluded as dead."),
        ("Primary test.", "van Elteren (1960) stratified Wilcoxon, exact by enumeration of within-stratum label assignments."),
        ("Effect size.", "Hedges' g with unequal-n variance, fixed-effect pooled across strata; DerSimonian–Laird random-effects g and I² alongside."),
        ("Multiplicity.", "Benjamini–Hochberg within each family of primary tests; per-stratum contrasts are secondary and uncorrected."),
        ("Outliers.", f"Modified z > {P.OUTLIER_Z} vs the group median, MAD pooled per run; both analyses reported."),
    ]:
        para(txt, bold_lead=lead, size=9)

    # ── Figures ──
    if figures:
        sec2 = doc.add_section(); sec2.orientation = WD_ORIENT.LANDSCAPE
        sec2.page_width, sec2.page_height = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec2, side, Cm(1.5))
        H("Figures")
        for cap, png in figures.items():
            doc.add_picture(io.BytesIO(png), width=Cm(26.0))
            para(cap, italic=True, size=8.5, colour=(0x50, 0x58, 0x64))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ─────────────────────────────────────────────────────────────────────────────
# Verify
# ─────────────────────────────────────────────────────────────────────────────
def verify(path: str, run: P.DamRun, run_b: Optional[P.DamRun] = None, cmp: Optional[dict] = None) -> Tuple[bool, str]:
    """
    Every numeric cell in the document's tables must be a display string the
    run's own tables produce. Returns (ok, human-readable report).
    """
    import docx
    runs = [run] + ([run_b] if run_b else [])
    pool = set()
    for r in runs:
        _, rows = group_rows(r)
        for row in rows:
            pool.update(row[1:])
        for fams in (r.contrasts, r.contrasts_noout):
            for df in fams.values():
                for c in contrast_rows(df):
                    pool.update([c["n"], c["ref"], c["cmp"], c["change"], c["g"], c["I2"], c["p"], c["q"], c["q"] + " **"])
                    # sensitivity table prints g and q on their own
                    pool.update([c["g"].split(" ")[0], c["q"] + " **"])
        for row in outlier_rows(r):
            pool.update(row)
        q = r.quality
        pool.update([fmt(q["bin_min"], 0), str(q["n_usable_days"]), f"{100*q['agreement']:.1f}%", str(len(r.per_fly)),
                     str(int(r.per_fly["Outlier"].sum())), f"{r.photoperiod_min/60:.1f}L:{24 - r.photoperiod_min/60:.1f}D",
                     f"{r.window[0].date()} to {(r.window[1] - pd.Timedelta(1, "D")).date()}" if r.window else "—",
                     str(q["cycling"])])
        s = r.summary()
        for v in s["families"].values():
            pool.update([str(v["n_tests"]), str(v["n_significant"])])
        dt = P.design_table(r.per_fly, r.design)
        pool.update(dt.astype(str).values.ravel().tolist())
    if cmp is not None:
        for c in contrast_rows(cmp["timepoint"]):
            pool.update([c["n"], c["ref"], c["cmp"], c["change"], c["g"], c["I2"], c["p"], c["q"], c["q"] + " **"])
        p = cmp["timepoint"]; p = p[p["analysis"] == "primary"]
        pool.update([str(len(p)), str(int(p["significant"].sum()))])
    numeric = re.compile(r"\d")
    doc = docx.Document(path)
    n_tot = n_bad = 0; bad = []
    for tb in doc.tables:
        for row in tb.rows[1:]:
            for c in row.cells[1:]:
                s = c.text.strip()
                if not numeric.search(s) or s == "—":
                    continue
                n_tot += 1
                if s not in pool:
                    n_bad += 1; bad.append(s)
    ok = n_bad == 0
    rep = (f"{n_tot - n_bad} of {n_tot} numeric table cells trace to the analysis tables"
           + ("" if ok else f"; unmatched: {bad[:8]}"))
    return ok, rep
